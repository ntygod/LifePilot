#!/usr/bin/env python3
"""Excel / CSV 常用处理工具。"""

from __future__ import annotations

import argparse
import csv
import json
import sys
from pathlib import Path
from typing import Any

try:
    from openpyxl import Workbook, load_workbook
except ImportError:
    print("错误：缺少 openpyxl。安装：python -m pip install openpyxl", file=sys.stderr)
    sys.exit(1)


def coerce(value: str, infer_types: bool) -> Any:
    if not infer_types:
        return value
    text = value.strip()
    if text == "":
        return None
    lowered = text.lower()
    if lowered in {"true", "false"}:
        return lowered == "true"
    try:
        if "." not in text:
            return int(text)
        return float(text)
    except ValueError:
        return value


def autosize(ws) -> None:
    for column in ws.columns:
        letter = column[0].column_letter
        width = min(60, max(10, max(len(str(cell.value or "")) for cell in column) + 2))
        ws.column_dimensions[letter].width = width


def from_csv(input_path: Path, output_path: Path, sheet_name: str, delimiter: str, infer_types: bool) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    with input_path.open("r", encoding="utf-8-sig", newline="") as fh:
        reader = csv.reader(fh, delimiter=delimiter)
        for row in reader:
            ws.append([coerce(value, infer_types) for value in row])
    autosize(ws)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f"已生成 {output_path}，sheet={sheet_name}, rows={ws.max_row}, cols={ws.max_column}")


def to_csv(input_path: Path, output_path: Path, sheet_name: str | None, delimiter: str) -> None:
    wb = load_workbook(input_path, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb.active
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh, delimiter=delimiter)
        for row in ws.iter_rows(values_only=True):
            writer.writerow(["" if value is None else value for value in row])
    print(f"已导出 {input_path}:{ws.title} -> {output_path}")


def inspect(input_path: Path) -> None:
    wb = load_workbook(input_path, read_only=True, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        sheets.append({
            "name": ws.title,
            "maxRow": ws.max_row,
            "maxColumn": ws.max_column,
        })
    print(json.dumps({"file": str(input_path), "sheets": sheets}, ensure_ascii=False, indent=2))


def main() -> int:
    parser = argparse.ArgumentParser(description="Excel / CSV 工具")
    sub = parser.add_subparsers(dest="command", required=True)

    p_from = sub.add_parser("from-csv", help="CSV 转 xlsx")
    p_from.add_argument("input")
    p_from.add_argument("output")
    p_from.add_argument("--sheet", default="Sheet1")
    p_from.add_argument("--delimiter", default=",")
    p_from.add_argument("--infer-types", action="store_true")

    p_to = sub.add_parser("to-csv", help="xlsx 转 CSV")
    p_to.add_argument("input")
    p_to.add_argument("output")
    p_to.add_argument("--sheet")
    p_to.add_argument("--delimiter", default=",")

    p_inspect = sub.add_parser("inspect", help="查看 workbook 概要")
    p_inspect.add_argument("input")

    args = parser.parse_args()
    try:
        if args.command == "from-csv":
            from_csv(Path(args.input), Path(args.output), args.sheet, args.delimiter, args.infer_types)
        elif args.command == "to-csv":
            to_csv(Path(args.input), Path(args.output), args.sheet, args.delimiter)
        elif args.command == "inspect":
            inspect(Path(args.input))
        return 0
    except Exception as exc:
        print(f"错误：{exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())

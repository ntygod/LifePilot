#!/usr/bin/env python3
"""Markdown 转 DOCX 的轻量脚本。

覆盖常用结构：标题、段落、无序/有序列表、代码块、简单 pipe 表格、图片引用。
复杂排版请优先用 pandoc 或直接写 python-docx 脚本。
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt
except ImportError:
    print("错误：缺少 python-docx。安装：python -m pip install python-docx", file=sys.stderr)
    sys.exit(1)


HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
UL_RE = re.compile(r"^\s*[-*+]\s+(.+)$")
OL_RE = re.compile(r"^\s*\d+[.)]\s+(.+)$")
IMAGE_RE = re.compile(r"!\[[^\]]*]\(([^)]+)\)")
TABLE_SEPARATOR_RE = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$")


def set_default_font(doc: Document, font_name: str) -> None:
    style = doc.styles["Normal"]
    style.font.name = font_name
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def add_text(doc: Document, text: str, style: str | None = None) -> None:
    text = text.strip()
    if not text:
        return
    paragraph = doc.add_paragraph(style=style)
    paragraph.add_run(text)


def add_code_block(doc: Document, lines: list[str]) -> None:
    if not lines:
        return
    paragraph = doc.add_paragraph()
    run = paragraph.add_run("\n".join(lines))
    run.font.name = "Courier New"
    run.font.size = Pt(9)
    paragraph.style = doc.styles["Normal"]


def split_table_row(line: str) -> list[str]:
    stripped = line.strip()
    if stripped.startswith("|"):
        stripped = stripped[1:]
    if stripped.endswith("|"):
        stripped = stripped[:-1]
    return [cell.strip() for cell in stripped.split("|")]


def looks_like_table(lines: list[str]) -> bool:
    return len(lines) >= 2 and TABLE_SEPARATOR_RE.match(lines[1]) is not None


def add_table(doc: Document, lines: list[str]) -> None:
    rows = [split_table_row(line) for i, line in enumerate(lines) if i != 1]
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        for c_idx in range(column_count):
            cell = table.cell(r_idx, c_idx)
            cell.text = row[c_idx] if c_idx < len(row) else ""
            if r_idx == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True


def add_image(doc: Document, md_path: Path, text: str, width_inches: float) -> bool:
    match = IMAGE_RE.search(text)
    if not match:
        return False
    raw_path = match.group(1).strip().strip("\"'")
    image_path = Path(raw_path)
    if not image_path.is_absolute():
        image_path = md_path.parent / image_path
    if not image_path.exists():
        add_text(doc, text)
        return True
    doc.add_picture(str(image_path), width=Inches(width_inches))
    caption = IMAGE_RE.sub("", text).strip()
    if caption:
        add_text(doc, caption)
    return True


def flush_paragraph(doc: Document, paragraph_lines: list[str]) -> None:
    if paragraph_lines:
        add_text(doc, " ".join(line.strip() for line in paragraph_lines))
        paragraph_lines.clear()


def convert(md_path: Path, docx_path: Path, font_name: str, image_width: float) -> None:
    if not md_path.exists():
        raise FileNotFoundError(f"输入文件不存在：{md_path}")

    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    set_default_font(doc, font_name)

    paragraph_lines: list[str] = []
    code_lines: list[str] = []
    table_lines: list[str] = []
    in_code = False

    def flush_table() -> None:
        nonlocal table_lines
        if not table_lines:
            return
        if looks_like_table(table_lines):
            flush_paragraph(doc, paragraph_lines)
            add_table(doc, table_lines)
        else:
            paragraph_lines.extend(table_lines)
        table_lines = []

    for line in lines:
        stripped = line.rstrip()

        if stripped.startswith("```"):
            flush_table()
            flush_paragraph(doc, paragraph_lines)
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue

        if in_code:
            code_lines.append(stripped)
            continue

        if "|" in stripped and stripped.strip().startswith("|"):
            table_lines.append(stripped)
            continue

        flush_table()

        if not stripped.strip():
            flush_paragraph(doc, paragraph_lines)
            continue

        heading = HEADING_RE.match(stripped)
        if heading:
            flush_paragraph(doc, paragraph_lines)
            level = min(len(heading.group(1)), 6)
            add_text(doc, heading.group(2), style=f"Heading {level}")
            continue

        ul = UL_RE.match(stripped)
        if ul:
            flush_paragraph(doc, paragraph_lines)
            add_text(doc, ul.group(1), style="List Bullet")
            continue

        ol = OL_RE.match(stripped)
        if ol:
            flush_paragraph(doc, paragraph_lines)
            add_text(doc, ol.group(1), style="List Number")
            continue

        if add_image(doc, md_path, stripped, image_width):
            flush_paragraph(doc, paragraph_lines)
            continue

        paragraph_lines.append(stripped)

    flush_table()
    flush_paragraph(doc, paragraph_lines)
    if in_code:
        add_code_block(doc, code_lines)

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(docx_path)
    print(f"已生成 {docx_path} ({docx_path.stat().st_size} bytes)")


def main() -> int:
    parser = argparse.ArgumentParser(description="Markdown 转 DOCX")
    parser.add_argument("input", help="输入 Markdown 文件")
    parser.add_argument("output", nargs="?", help="输出 DOCX 文件；默认同名 .docx")
    parser.add_argument("--font", default="Microsoft YaHei", help="默认字体")
    parser.add_argument("--image-width", type=float, default=5.5, help="图片宽度（英寸）")
    args = parser.parse_args()

    md_path = Path(args.input)
    output = Path(args.output) if args.output else md_path.with_suffix(".docx")
    try:
        convert(md_path, output, args.font, args.image_width)
        return 0
    except Exception as exc:
        print(f"错误：{exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
